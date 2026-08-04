# -*- coding: utf-8 -*-
"""مكتبة بناء العروض التقديمية العربية RTL — نظام أثير"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

NAVY  = RGBColor(0x12, 0x31, 0x4E)
NAVY2 = RGBColor(0x1D, 0x4E, 0x79)
NAVY3 = RGBColor(0x2A, 0x5A, 0x87)
GOLD  = RGBColor(0xC9, 0xA2, 0x27)
TEAL  = RGBColor(0x12, 0x87, 0x71)
INK   = RGBColor(0x1F, 0x29, 0x33)
GRAY  = RGBColor(0x5B, 0x64, 0x70)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
LIGHT2= RGBColor(0xED, 0xF3, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED   = RGBColor(0xB0, 0x3A, 0x2E)

SLD_W = Inches(13.333)
SLD_H = Inches(7.5)
FONT = 'Segoe UI'

def new_deck():
    prs = Presentation()
    prs.slide_width = SLD_W
    prs.slide_height = SLD_H
    return prs

def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def _set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set('rtl', '1')
    pPr.set('algn', 'r')

def _set_font(run, size, bold, color, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:cs'):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set('typeface', FONT)

def add_text(slide, text, x, y, w, h, size=16, bold=False, color=INK,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15,
             wrap=True, italic=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if align == PP_ALIGN.RIGHT:
            _set_rtl(p)
        r = p.add_run()
        r.text = line
        _set_font(r, size, bold, color, italic)
    return box

def add_bullets(slide, items, x, y, w, h, size=15, color=INK, gap=8,
                marker='•', bold_first=False, marker_color=GOLD):
    """نقاط مرقمة/منقطة RTL — العلامة تلحق بداية النص (يمين)."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    n = 0
    for it in items:
        n += 1
        if isinstance(it, tuple):
            head, rest = it
        else:
            head, rest = None, it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.RIGHT
        p.line_spacing = 1.18
        p.space_after = Pt(gap)
        _set_rtl(p)
        mk = marker if marker != '#' else f'{n}.'
        r0 = p.add_run(); r0.text = f'{mk}  '
        _set_font(r0, size, True, marker_color)
        if head:
            r1 = p.add_run(); r1.text = head + ' '
            _set_font(r1, size, True, NAVY2)
        r2 = p.add_run(); r2.text = rest
        _set_font(r2, size, bold_first and bool(head), color)
    return box

def rect(slide, x, y, w, h, fill=NAVY, line=None, line_w=None,
         shape=MSO_SHAPE.RECTANGLE, shadow_off=True):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = line_w or Pt(1)
    if shadow_off:
        s.shadow.inherit = False
    return s

def title_bar(slide, title, kicker=None, num=None):
    rect(slide, 0, 0, SLD_W, Inches(1.05), fill=NAVY)
    rect(slide, 0, Inches(1.05), SLD_W, Inches(0.045), fill=GOLD)
    add_text(slide, title, Inches(0.5), Inches(0.06), Inches(11.2),
             Inches(0.62), size=25, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        add_text(slide, kicker, Inches(0.5), Inches(0.66), Inches(11.2),
                 Inches(0.35), size=11, bold=False, color=RGBColor(0xBC, 0xCB, 0xDB))
    if num:
        circ = rect(slide, Inches(0.35), Inches(0.22), Inches(0.62), Inches(0.62),
                    fill=GOLD, shape=MSO_SHAPE.OVAL)
        tf = circ.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(num)
        _set_font(r, 20, True, NAVY)

def footer(slide, text='أثير لإدارة الفنادق — وثيقة تقديمية'):
    add_text(slide, text, Inches(0.4), SLD_H - Inches(0.42), Inches(12.5),
             Inches(0.35), size=9, color=GRAY, align=PP_ALIGN.LEFT)

def cover_slide(prs, title, subtitle, lines, badge=None):
    s = blank(prs)
    rect(s, 0, 0, SLD_W, SLD_H, fill=NAVY)
    rect(s, 0, 0, SLD_W, Inches(0.18), fill=GOLD)
    rect(s, 0, SLD_H - Inches(0.18), SLD_W, Inches(0.18), fill=GOLD)
    rect(s, Inches(11.2), Inches(1.6), Inches(0.09), Inches(3.4), fill=GOLD)
    add_text(s, title, Inches(0.8), Inches(1.5), Inches(10.2), Inches(1.6),
             size=44, bold=True, color=WHITE)
    add_text(s, subtitle, Inches(0.8), Inches(2.85), Inches(10.2), Inches(0.7),
             size=20, bold=True, color=GOLD)
    add_bullets(s, lines, Inches(0.8), Inches(3.8), Inches(10.0), Inches(2.8),
                size=15, color=RGBColor(0xD5, 0xDE, 0xE8), gap=10,
                marker='—', marker_color=GOLD)
    if badge:
        b = rect(s, Inches(8.6), SLD_H - Inches(1.15), Inches(4.0), Inches(0.55),
                 fill=NAVY2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = b.text_frame; tf.word_wrap = False
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _set_rtl(p)
        r = p.add_run(); r.text = badge
        _set_font(r, 13, True, GOLD)
    return s

def card(slide, x, y, w, h, title, body, fill=LIGHT, tcolor=NAVY, tsize=15,
         bsize=12.5, bar=GOLD, body_color=INK):
    rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, x, y, w, Inches(0.09), fill=bar)
    add_text(slide, title, x + Inches(0.12), y + Inches(0.13), w - Inches(0.24),
             Inches(0.55), size=tsize, bold=True, color=tcolor)
    add_text(slide, body, x + Inches(0.12), y + Inches(0.62), w - Inches(0.24),
             h - Inches(0.75), size=bsize, color=body_color, line_spacing=1.12)

def chevron_flow(slide, steps, y, x_end=SLD_W - Inches(0.5), h=Inches(0.85),
                 colors=None, size=13):
    """أسهم انسيابية يمين→يسار"""
    n = len(steps)
    gap = Inches(0.08)
    total_w = x_end - Inches(0.5)
    cw = Emu(int((total_w - gap * (n - 1)) / n))
    colors = colors or [NAVY2, TEAL]
    for i, stext in enumerate(steps):
        x = Emu(int(x_end) - (i + 1) * int(cw) - i * int(gap))
        shp = MSO_SHAPE.CHEVRON if i < n - 1 else MSO_SHAPE.PENTAGON
        c = rect(slide, x, y, cw, h, fill=colors[i % 2], shape=shp)
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _set_rtl(p)
        r = p.add_run(); r.text = stext
        _set_font(r, size, True, WHITE)
        c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

def big_quote(slide, text, y, h=Inches(1.5), sub=None):
    rect(slide, Inches(0.7), y, SLD_W - Inches(1.4), h, fill=NAVY,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, text, Inches(1.1), y + Inches(0.15), SLD_W - Inches(2.2),
             h - Inches(0.3), size=22, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    return

def stat_big(slide, x, y, w, h, number, label, ncolor=NAVY2):
    rect(slide, x, y, w, h, fill=LIGHT2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, number, x, y + Inches(0.18), w, Inches(0.75), size=32,
             bold=True, color=ncolor, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + Inches(0.1), y + Inches(0.95), w - Inches(0.2),
             h - Inches(1.05), size=12, color=GRAY, align=PP_ALIGN.CENTER)

def table_slide(slide, headers, rows, x, y, w, h, col_w=None, hfs=13,
                rfs=11.5, right_cols=None):
    """جدول RTL: الأعمدة منطقياً أولها يمين"""
    nrows = len(rows) + 1
    ncols = len(headers)
    gt = slide.shapes.add_table(nrows, ncols, x, y, w, h).table
    gt.horz_banding = True
    if col_w:
        total = sum(col_w)
        for i, cwv in enumerate(reversed(col_w)):
            gt.columns[i].width = Emu(int(int(w) * cwv / total))
    for i, htext in enumerate(reversed(headers)):
        c = gt.cell(0, i)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.margin_top = c.margin_bottom = Inches(0.03)
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER; _set_rtl(p)
        r = p.add_run(); r.text = htext
        _set_font(r, hfs, True, WHITE)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(reversed(row)):
            c = gt.cell(ri, ci)
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if ri % 2 else LIGHT
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER; _set_rtl(p)
            r = p.add_run(); r.text = str(val)
            _set_font(r, rfs, False, INK)
    return gt
